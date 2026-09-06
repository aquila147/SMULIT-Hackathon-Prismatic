import fitz
import pytesseract
from PIL import Image
import io
import os
import logging
import re
from pathlib import Path

logger = logging.getLogger(__name__)

SCANNED_THRESHOLD = 50

PAGES_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "page_images")
os.makedirs(PAGES_DIR, exist_ok=True)


def ingest_file(filepath: str) -> dict:
    """
    Main entry point. Detects file type and routes to the correct handler.
    Returns the same shape regardless of input type.
    """
    ext = Path(filepath).suffix.lower()

    if ext == ".pdf":
        return ingest_pdf(filepath)
    elif ext == ".docx":
        return ingest_docx(filepath)
    elif ext in (".pptx", ".ppt"):
        return ingest_pptx(filepath)
    elif ext == ".txt":
        return ingest_txt(filepath)
    else:
        raise ValueError(f"Unsupported file type: {ext}. Supported formats: PDF, DOCX, PPTX, TXT")


def ingest_pdf(filepath: str) -> dict:
    """
    Extract text and coordinates from every page of a PDF.
    Handles scanned pages via OCR.
    Raises ValueError for password protected or corrupt files.
    """
    # Try opening the file
    try:
        doc = fitz.open(filepath)
    except Exception as e:
        raise ValueError(f"Could not open PDF. The file may be corrupt: {filepath}")

    # Check for password protection
    if doc.is_encrypted:
        doc.close()
        raise ValueError(
            "PASSWORD_PROTECTED: This PDF is password protected. "
            "Please remove the password using Adobe Acrobat or smallpdf.com before uploading."
        )

    filename = os.path.basename(filepath)
    pages = []
    full_text_parts = []
    any_scanned = False

    for page_num in range(len(doc)):
        page = doc[page_num]

        # Extract text blocks first
        text_blocks = []
        page_text_parts = []

        for block in page.get_text("dict")["blocks"]:
            if block["type"] == 0:
                block_text = ""
                for line in block["lines"]:
                    for span in line["spans"]:
                        block_text += span["text"]
                    block_text += "\n"
                block_text = block_text.strip()
                if block_text:
                    text_blocks.append({
                        "text": block_text,
                        "bbox": list(block["bbox"]),
                        "page": page_num,
                    })
                    page_text_parts.append(block_text)

        page_text = "\n".join(page_text_parts)

        is_ocr = False
        image_path = None

        if len(page_text.strip()) < SCANNED_THRESHOLD:
            # Scanned page — must rasterise for OCR
            logger.info(f"Page {page_num}: scanned, running OCR")
            any_scanned = True
            is_ocr = True
            try:
                page_text, text_blocks, image_path = _ocr_page(page, page_num, filename)
            except Exception as e:
                logger.warning(f"OCR failed on page {page_num}: {e}")
        else:
            # Text-based page — no rasterisation needed
            logger.info(f"Page {page_num}: text-based, {len(page_text)} chars")

        pages.append({
            "page_number": page_num,
            "text": page_text,
            "blocks": text_blocks,
            "is_ocr": is_ocr,
            "image_path": image_path,
        })
        full_text_parts.append(page_text)

    doc.close()
    full_text = "\n\n".join(full_text_parts)

    return {
        "filename": filename,
        "filepath": filepath,
        "file_type": "pdf",
        "is_scanned": any_scanned,
        "pages": pages,
        "full_text": full_text,
        "clauses_total": _estimate_clause_count(full_text),
    }


def ingest_docx(filepath: str) -> dict:
    """
    Extract text from a Word document.
    No bboxes — Word has no page coordinate system.
    Click-to-citation will not be available for these files.
    """
    try:
        from docx import Document
    except ImportError:
        raise ValueError("python-docx not installed. Run: pip install python-docx")

    try:
        doc = Document(filepath)
    except Exception as e:
        raise ValueError(f"Could not open Word document. The file may be corrupt: {filepath}")

    filename = os.path.basename(filepath)

    # Extract all paragraph text
    paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]

    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())

    full_text = "\n".join(paragraphs)

    return {
        "filename": filename,
        "filepath": filepath,
        "file_type": "docx",
        "is_scanned": False,
        "pages": [{
            "page_number": 0,
            "text": full_text,
            "blocks": [],       # No bboxes available for Word docs
            "is_ocr": False,
            "image_path": None,
        }],
        "full_text": full_text,
        "clauses_total": _estimate_clause_count(full_text),
        "warning": "Word document — clause highlighting not available. Text extraction only."
    }


def ingest_pptx(filepath: str) -> dict:
    """
    Extract text from a PowerPoint presentation.
    Each slide becomes a 'page'.
    No bboxes — click-to-citation not available.
    """
    try:
        from pptx import Presentation
    except ImportError:
        raise ValueError("python-pptx not installed. Run: pip install python-pptx")

    try:
        prs = Presentation(filepath)
    except Exception as e:
        raise ValueError(f"Could not open PowerPoint file. The file may be corrupt: {filepath}")

    filename = os.path.basename(filepath)
    pages = []
    full_text_parts = []

    for slide_num, slide in enumerate(prs.slides):
        # Extract text from all shapes on the slide
        slide_texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_texts.append(shape.text.strip())

        slide_text = "\n".join(slide_texts)

        pages.append({
            "page_number": slide_num,
            "text": slide_text,
            "blocks": [],       # No bboxes available for PowerPoint
            "is_ocr": False,
            "image_path": None,
        })
        full_text_parts.append(slide_text)

    full_text = "\n\n".join(full_text_parts)

    return {
        "filename": filename,
        "filepath": filepath,
        "file_type": "pptx",
        "is_scanned": False,
        "pages": pages,
        "full_text": full_text,
        "clauses_total": _estimate_clause_count(full_text),
        "warning": "PowerPoint file — clause highlighting not available. Text extraction only."
    }


def ingest_txt(filepath: str) -> dict:
    """
    Extract text from a plain text file.
    No bboxes — click-to-citation not available.
    """
    try:
        with open(filepath, "r", encoding="utf-8") as f:
            text = f.read()
    except UnicodeDecodeError:
        # Fallback to latin-1 if UTF-8 fails
        with open(filepath, "r", encoding="latin-1") as f:
            text = f.read()

    filename = os.path.basename(filepath)

    return {
        "filename": filename,
        "filepath": filepath,
        "file_type": "txt",
        "is_scanned": False,
        "pages": [{
            "page_number": 0,
            "text": text,
            "blocks": [],
            "is_ocr": False,
            "image_path": None,
        }],
        "full_text": text,
        "clauses_total": _estimate_clause_count(text),
        "warning": "Text file — clause highlighting not available. Text extraction only."
    }


def _ocr_page(page, page_num: int, filename: str) -> tuple:
    """
    Only called for scanned PDF pages.
    Rasterises at 300 DPI for OCR, saves 150 DPI copy for the viewer.
    """
    mat_ocr = fitz.Matrix(300 / 72, 300 / 72)
    pix = page.get_pixmap(matrix=mat_ocr)
    img_bytes = pix.tobytes("png")
    img = Image.open(io.BytesIO(img_bytes))

    ocr_data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT)

    blocks = []
    current_block_text = []
    current_block_bbox = None

    for i in range(len(ocr_data["text"])):
        word = ocr_data["text"][i].strip()
        conf = int(ocr_data["conf"][i])

        if not word or conf < 30:
            if current_block_text:
                blocks.append({
                    "text": " ".join(current_block_text),
                    "bbox": current_block_bbox,
                    "page": page_num,
                })
                current_block_text = []
                current_block_bbox = None
            continue

        scale = 72 / 300
        x0 = ocr_data["left"][i] * scale
        y0 = ocr_data["top"][i] * scale
        x1 = (ocr_data["left"][i] + ocr_data["width"][i]) * scale
        y1 = (ocr_data["top"][i] + ocr_data["height"][i]) * scale

        current_block_text.append(word)
        if current_block_bbox is None:
            current_block_bbox = [x0, y0, x1, y1]
        else:
            current_block_bbox[0] = min(current_block_bbox[0], x0)
            current_block_bbox[1] = min(current_block_bbox[1], y0)
            current_block_bbox[2] = max(current_block_bbox[2], x1)
            current_block_bbox[3] = max(current_block_bbox[3], y1)

    if current_block_text:
        blocks.append({
            "text": " ".join(current_block_text),
            "bbox": current_block_bbox,
            "page": page_num,
        })

    full_text = " ".join(b["text"] for b in blocks)

    # Save 150 DPI viewer image for scanned pages only
    mat_view = fitz.Matrix(150 / 72, 150 / 72)
    pix_view = page.get_pixmap(matrix=mat_view)
    safe_name = Path(filename).stem
    image_path = os.path.join(PAGES_DIR, f"{safe_name}_p{page_num}.png")
    pix_view.save(image_path)

    return full_text, blocks, image_path


def _estimate_clause_count(text: str) -> int:
    patterns = [
        r'^\s*\d+\.\d*\s',
        r'^\s*\(\w+\)\s',
        r'^\s*Section\s+\d+',
        r'^\s*Article\s+\d+',
        r'^\s*Clause\s+\d+',
    ]
    count = 0
    for line in text.split("\n"):
        for pat in patterns:
            if re.match(pat, line, re.IGNORECASE):
                count += 1
                break
    return max(count, 1)


def chunk_document(doc_data: dict, chunk_size: int = 4000, overlap: int = 300) -> list:
    # Raised from 2000/200: bigger, fewer chunks means fewer LLM round trips
    # for Agent 2 (one call per chunk) without meaningfully hurting extraction
    # quality — 4000 chars is still well inside Agent 2's 8192-token response
    # budget. Total call count is the real lever on both speed and on
    # avoiding backend rate/concurrency errors; this halves it for Agent 2
    # the same way batching did for Agents 3 and 5.
    full_text = doc_data["full_text"]
    chunks = []
    start = 0

    while start < len(full_text):
        end = min(start + chunk_size, len(full_text))

        if end < len(full_text):
            newline_pos = full_text.rfind("\n\n", start, end)
            if newline_pos > start + chunk_size // 2:
                end = newline_pos + 2

        chunk_text = full_text[start:end]

        pages_in_chunk = set()
        char_count = 0
        for page in doc_data["pages"]:
            page_len = len(page["text"]) + 2
            if char_count + page_len > start and char_count < end:
                pages_in_chunk.add(page["page_number"])
            char_count += page_len

        chunks.append({
            "text": chunk_text,
            "start_char": start,
            "end_char": end,
            "pages": sorted(pages_in_chunk),
        })

        # If we reached the end, stop immediately
        if end >= len(full_text):
            break

        start = end - overlap

    return chunks